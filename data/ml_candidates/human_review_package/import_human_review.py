"""
data/ml_candidates/human_review_package/import_human_review.py

Import and Validation Script for Human-Reviewed Labels.
Applies human review decisions from human_review_workbook.xlsx to create a
verified dataset output.

Safety Constraints:
1. NEVER overwrites or alters the canonical dataset (data/ml_candidates/final_dataset.jsonl).
2. Refuses to mark a label reviewed if decision is blank or 'Need help'.
3. Preserves unresolved decisions as 'unresolved'.
4. Requires reviewer name/initials and date for every applied human decision.
5. NEVER marks an entire record 'human_verified: true' if ANY of its 5 category labels
   remain unreviewed, pending, or unresolved.
6. Generates a detailed audit change report.
"""

import os
import sys
import json
import datetime
import openpyxl

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PACKAGE_DIR = SCRIPT_DIR
WORKBOOK_PATH = os.path.join(PACKAGE_DIR, "human_review_workbook.xlsx")
REVIEW_JSONL_PATH = os.path.join(PACKAGE_DIR, "human_review_records.jsonl")
CANONICAL_PATH = os.path.join(PACKAGE_DIR, "..", "final_dataset.jsonl")
OUTPUT_PATH = os.path.join(PACKAGE_DIR, "..", "human_reviewed_dataset.jsonl")
REPORT_PATH = os.path.join(PACKAGE_DIR, "human_review_change_report.txt")

VALID_DECISIONS = {"Confirm", "Correct", "Unresolved", "Need help"}
VALID_STATUSES = {"positive", "negative", "unresolved"}

def import_human_reviews():
    print(f"Starting human review import from: {WORKBOOK_PATH}")
    if not os.path.exists(WORKBOOK_PATH):
        print(f"ERROR: Workbook not found at {WORKBOOK_PATH}")
        sys.exit(1)

    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    if "CATEGORY LABEL REVIEW" not in wb.sheetnames:
        print("ERROR: Sheet 'CATEGORY LABEL REVIEW' not found in workbook.")
        sys.exit(1)

    ws = wb["CATEGORY LABEL REVIEW"]
    
    # Read review decisions
    # Columns:
    # A(1): Sample ID, B(2): Origin, C(3): Project ID, D(4): Contract/Function, E(5): Category Name
    # F(6): AI Status, G(7): AI Severity, H(8): AI Evidence, I(9): AI Rationale
    # P(16): Reviewer Decision, Q(17): Corrected Status, R(18): Corrected Evidence,
    # S(19): Reviewer Notes, T(20): Reviewer Name, U(21): Review Date
    
    decisions = {}  # (sample_id, category) -> dict of decision info
    row_count = 0
    completed_decisions = 0
    blank_decisions = 0
    need_help_decisions = 0
    validation_errors = []

    for row in range(2, ws.max_row + 1):
        sid = ws.cell(row=row, column=1).value
        cat = ws.cell(row=row, column=5).value
        ai_status = ws.cell(row=row, column=6).value
        ai_evidence = ws.cell(row=row, column=8).value
        
        decision = ws.cell(row=row, column=16).value
        corrected_status = ws.cell(row=row, column=17).value
        corrected_evidence = ws.cell(row=row, column=18).value
        notes = ws.cell(row=row, column=19).value
        reviewer = ws.cell(row=row, column=20).value
        rev_date = ws.cell(row=row, column=21).value

        if not sid or not cat:
            continue

        sid = str(sid).strip()
        cat = str(cat).strip()
        row_count += 1

        if not decision or str(decision).strip() == "":
            blank_decisions += 1
            decisions[(sid, cat)] = {"status": "unreviewed"}
            continue

        decision = str(decision).strip()
        if decision == "Need help":
            need_help_decisions += 1
            decisions[(sid, cat)] = {"status": "need_help", "notes": notes, "reviewer": reviewer}
            continue

        if decision not in VALID_DECISIONS:
            validation_errors.append(f"Row {row} ({sid} - {cat}): Invalid decision '{decision}'")
            continue

        # Require reviewer name and date for confirmed/corrected/unresolved
        if not reviewer or str(reviewer).strip() == "":
            validation_errors.append(f"Row {row} ({sid} - {cat}): Missing reviewer name/initials")
            continue

        if not rev_date:
            validation_errors.append(f"Row {row} ({sid} - {cat}): Missing review date")
            continue

        if decision == "Correct":
            if not corrected_status or str(corrected_status).strip() not in VALID_STATUSES:
                validation_errors.append(f"Row {row} ({sid} - {cat}): Decision is 'Correct' but 'Corrected Status' is missing or invalid ({corrected_status})")
                continue

        decisions[(sid, cat)] = {
            "status": "reviewed",
            "decision": decision,
            "ai_status": ai_status,
            "ai_evidence": ai_evidence,
            "corrected_status": str(corrected_status).strip() if corrected_status else None,
            "corrected_evidence": str(corrected_evidence).strip() if corrected_evidence else None,
            "notes": str(notes).strip() if notes else None,
            "reviewer": str(reviewer).strip(),
            "date": str(rev_date).strip()
        }
        completed_decisions += 1

    print(f"Total rows inspected: {row_count}")
    print(f"Completed human decisions: {completed_decisions}")
    print(f"Blank (unreviewed) decisions: {blank_decisions}")
    print(f"Need help decisions: {need_help_decisions}")

    if validation_errors:
        print(f"\nVALIDATION ERRORS FOUND ({len(validation_errors)}):")
        for err in validation_errors[:10]:
            print(f"  - {err}")
        if len(validation_errors) > 10:
            print(f"  ... and {len(validation_errors) - 10} more errors.")
        print("\nPromotion aborted due to validation errors. Please correct the workbook.")
        sys.exit(1)

    if completed_decisions == 0:
        print("\nNOTICE: Zero human decisions recorded. Workbook has not yet been filled out.")
        print("Promotion script exited safely without creating output.")
        sys.exit(0)

    # Load review copy JSONL
    with open(REVIEW_JSONL_PATH, "r", encoding="utf-8") as f:
        records = [json.loads(line) for line in f if line.strip()]

    output_records = []
    change_log = []
    fully_verified_records = 0

    for r in records:
        sid = r["sample_id"]
        record_categories_verified = 0

        for cat in ["Privacy", "Termination", "Legal Compliance", "Liability", "Payment / Financial"]:
            key = (sid, cat)
            if key in decisions and decisions[key]["status"] == "reviewed":
                d = decisions[key]
                dec_type = d["decision"]
                
                # Apply decision
                r["human_category_reviews"][cat]["decision"] = dec_type
                r["human_category_reviews"][cat]["reviewer"] = d["reviewer"]
                r["human_category_reviews"][cat]["review_date"] = d["date"]
                r["human_category_reviews"][cat]["reviewer_notes"] = d["notes"]

                if dec_type == "Confirm":
                    r["human_category_reviews"][cat]["reviewed_status"] = d["ai_status"]
                    r["human_category_reviews"][cat]["reviewed_evidence"] = d["ai_evidence"]
                    record_categories_verified += 1
                    change_log.append(f"[{sid} | {cat}] CONFIRMED as {d['ai_status']} by {d['reviewer']}")

                elif dec_type == "Correct":
                    old_st = d["ai_status"]
                    new_st = d["corrected_status"]
                    r["human_category_reviews"][cat]["reviewed_status"] = new_st
                    r["human_category_reviews"][cat]["reviewed_evidence"] = d["corrected_evidence"]
                    # Update active category label
                    r["categories"][cat]["status"] = new_st
                    if d["corrected_evidence"]:
                        r["categories"][cat]["evidence"] = d["corrected_evidence"]
                    r["categories"][cat]["rationale"] = f"Corrected by human reviewer ({d['reviewer']}): {d['notes']}"
                    record_categories_verified += 1
                    change_log.append(f"[{sid} | {cat}] CORRECTED from {old_st} to {new_st} by {d['reviewer']}")

                elif dec_type == "Unresolved":
                    r["human_category_reviews"][cat]["reviewed_status"] = "unresolved"
                    r["categories"][cat]["status"] = "unresolved"
                    r["categories"][cat]["rationale"] = f"Marked unresolved during human review ({d['reviewer']}): {d['notes']}"
                    change_log.append(f"[{sid} | {cat}] Marked UNRESOLVED by {d['reviewer']}")

        # Full record human verification check
        if record_categories_verified == 5:
            r["human_verified"] = True
            r["provenance"]["human_verified"] = True
            r["provenance"]["human_review_status"] = "human_verified"
            r["provenance"]["human_reviewer"] = d["reviewer"]
            r["provenance"]["human_review_date"] = d["date"]
            r["review_status"] = "human_verified"
            fully_verified_records += 1
        elif record_categories_verified > 0:
            r["human_verified"] = False
            r["provenance"]["human_verified"] = False
            r["provenance"]["human_review_status"] = "partial_human_review"
            r["review_status"] = "ai_reviewed"

        output_records.append(r)

    # Write output to SEPARATE file
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        for r in output_records:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    # Write change report
    with open(REPORT_PATH, "w", encoding="utf-8") as f:
        f.write("Smart Contract Analyser — Human Review Promotion Change Report\n")
        f.write(f"Generated: {datetime.datetime.now().isoformat()}\n")
        f.write(f"Input Workbook: {WORKBOOK_PATH}\n")
        f.write(f"Output Dataset: {OUTPUT_PATH}\n")
        f.write(f"Canonical Dataset Untouched: {CANONICAL_PATH}\n\n")
        f.write(f"Summary:\n")
        f.write(f"  Total records processed: {len(output_records)}\n")
        f.write(f"  Fully human-verified records (5/5 categories confirmed/corrected): {fully_verified_records}\n")
        f.write(f"  Individual category decisions applied: {len(change_log)}\n\n")
        f.write("Detailed Decision Log:\n")
        for line in change_log:
            f.write(f"  - {line}\n")

    print(f"\nSUCCESS: Human review imported successfully.")
    print(f"Output saved to: {OUTPUT_PATH}")
    print(f"Change report saved to: {REPORT_PATH}")
    print(f"Canonical dataset ({CANONICAL_PATH}) remains completely untouched.")

if __name__ == "__main__":
    import_human_reviews()
